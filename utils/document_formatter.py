import re, tempfile, os
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from fpdf import FPDF

def sanitize_text(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("```text", "").replace("```markdown", "").replace("```", "").strip()

def format_docx(text, doc_type, company_name, logo_bytes=None):
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.7); sec.left_margin = sec.right_margin = Inches(.8)
    if logo_bytes:
        try:
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(BytesIO(logo_bytes), width=Inches(1.15))
        except Exception: pass
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(doc_type.upper()); r.bold=True; r.font.name="Times New Roman"; r.font.size=Pt(16)
    if company_name:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(company_name); r.font.size=Pt(10)
    for line in sanitize_text(text).splitlines():
        s=line.strip()
        if not s: continue
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5)
        r=p.add_run(s); r.font.name="Times New Roman"; r.font.size=Pt(11)
        if re.match(r"^(SECTION\s+\d+|ARTICLE\s+[IVX\d]+|\d+[.)]\s+|[A-Z][A-Z\s/&-]{5,}:)$", s): r.bold=True
    table=doc.add_table(rows=1, cols=2); table.style="Table Grid"; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text="No."; table.rows[0].cells[1].text="Key Term"
    for i, term in enumerate([x.strip() for x in re.split(r";|\n", text) if x.strip()][:12],1):
        cells=table.add_row().cells; cells[0].text=str(i); cells[1].text=term[:500]
    footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase • AI-assisted draft • Review with a qualified legal professional before use.")
    out=BytesIO(); doc.save(out); return out.getvalue()

class LegalPDF(FPDF):
    def footer(self):
        self.set_y(-15); self.set_font("Helvetica", size=8)
        self.cell(0,8,"LegalEase • AI-assisted draft • Review with a qualified legal professional before use.",align="C")

def format_pdf(text, doc_type, company_name, logo_bytes=None):
    pdf=LegalPDF(); pdf.set_auto_page_break(auto=True, margin=18); pdf.add_page()
    if logo_bytes:
        path=None
        try:
            suffix=".png"
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f: f.write(logo_bytes); path=f.name
            pdf.image(path, x=88, y=12, w=28); pdf.ln(24)
        except Exception: pdf.ln(4)
        finally:
            if path:
                try: os.unlink(path)
                except OSError: pass
    pdf.set_font("Helvetica","B",16); pdf.cell(0,10,doc_type.upper(),new_x="LMARGIN",new_y="NEXT",align="C")
    if company_name:
        pdf.set_font("Helvetica",size=9); pdf.cell(0,6,company_name,new_x="LMARGIN",new_y="NEXT",align="C")
    pdf.ln(4); pdf.set_font("Helvetica",size=10)
    for line in sanitize_text(text).splitlines():
        line=line.strip()
        if not line: pdf.ln(3); continue
        if re.match(r"^(SECTION\s+\d+|ARTICLE\s+[IVX\d]+|\d+[.)]\s+|[A-Z][A-Z\s/&-]{5,}:)$",line): pdf.set_font("Helvetica","B",11)
        else: pdf.set_font("Helvetica",size=10)
        pdf.multi_cell(0,5.5,line)
    return bytes(pdf.output(dest="S"))
